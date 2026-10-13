from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 默认中文字体
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

BLUE = RGBColor(0x1d, 0x4e, 0x89)

# 标题
doc.add_heading('猎聘网简历自动初筛工具', level=0)
sub = doc.add_paragraph('HR 招聘初筛提效方案 —— 扫描 · 解析 · AI 评估 · 复核')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic = True
sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# 一、工具定位
doc.add_heading('一、工具定位', level=1)
doc.add_paragraph(
    '将招聘中“逐份人工打开简历、阅读、判断”的初筛环节，升级为'
    '系统自动扫描解析 + 大模型 AI 深度评估 + 人工复核确认。'
    '既保留 HR 的最终决策权，又把大量重复阅读与基础判断交给工具完成，'
    '显著缩短初筛周期、统一筛选标准。'
)

# 二、核心流程
doc.add_heading('二、核心流程', level=1)
flows = [
    ('一键启动', '后台服务与浏览器环境自动就绪'),
    ('自动扫描', '读取猎聘“收到简历 / 未处理”列表'),
    ('智能解析', '提取姓名、目前公司/岗位、期望城市、工作经历、应聘岗位'),
    ('AI 评估', '规则引擎 + 大模型双重评估，输出匹配摘要与关注点'),
    ('自动判定', '合适→通过筛选 / 不合适→标记不合适（可仅模拟）'),
    ('看板复核', 'HR 统一查看、跳转详情、确认操作'),
]
for title, desc in flows:
    p = doc.add_paragraph(style='List Number')
    r = p.add_run(title + '：')
    r.bold = True
    p.add_run(desc)

# 三、主要功能
doc.add_heading('三、主要功能', level=1)
funcs = [
    ('自动扫描', '定时/手动读取猎聘未处理简历，无需逐份手动点开。'),
    ('精准解析', '目前公司/岗位、期望城市、应聘岗位与岗位 JD 自动匹配。'),
    ('大模型深度评估', '调用 MiniMax 大模型生成人岗匹配摘要与风险关注点。'),
    ('双引擎兜底', '规则版与 LLM 互备，大模型异常时自动回退规则版，稳定可靠。'),
    ('看板管理', '候选人、时间、AI 评估一目了然，支持跳转猎聘详情页。'),
    ('安全可控', '默认“仅模拟”预览，真实点击操作需人工确认，避免误处理。'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
table.rows[0].cells[0].text = '功能'
table.rows[0].cells[1].text = '说明'
for name, desc in funcs:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = desc

# 四、效率提升亮点
doc.add_heading('四、效率提升亮点', level=1)
kpis = doc.add_paragraph()
for label in ['80%+  预计节省初筛工时    ', '秒级  单份简历评估耗时    ', '7 类  岗位 JD 已覆盖']:
    r = kpis.add_run(label)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = BLUE

tips = [
    '初筛从“人工逐份阅读 + 判断”变为“系统秒级自动评估 + 人工复核”，HR 只需聚焦系统标注的待定/异常项。',
    '大模型替代人工全文阅读，判断更客观、标准更统一，减少主观偏差与漏看。',
    '支持批量扫描与每日增量处理，可稳定支撑招聘高峰期的常态化初筛。',
    '风险关注点自动高亮（如经验缺口、资质不符），辅助 HR 快速决策。',
]
for tip in tips:
    doc.add_paragraph(tip, style='List Bullet')

# 当前状态
doc.add_heading('当前状态与扩展', level=1)
doc.add_paragraph(
    '已在本地稳定运行，覆盖 CRA、SSU、临床项目经理、医学经理等多类岗位 JD。'
    '后续可扩展：与飞书多维表格打通候选人库、支持更多招聘平台接入。'
)

out = r'D:\腾讯小龙虾\处理简历\猎聘网处理简历\简历初筛工具简介.docx'
doc.save(out)
print('saved', out)
